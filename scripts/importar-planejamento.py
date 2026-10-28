"""
Importa o planejamento do milho das planilhas para o Supabase.

Lê as abas PLANEJAMENTO de MILHO_2025 e MILHO_2026 e a SOCIEDADE GILBERTO
(MILHO_2025) e gera supabase/importacao-planejamento.sql, para colar no SQL
Editor.

    python3 scripts/importar-planejamento.py <pasta com os .xlsx>

Cada bloco "FAZENDA X" da planilha vira os itens da safra daquela fazenda:
um item por produto de cada fase (dose/ha, quantidade, preço, total) e um item
para cada custo do resumo à direita (operacional, colheita, arrendamento,
juros), com o valor por hectare no preço. Os produtos que ainda não estão no
cadastro são criados (comparando o nome sem acento); o mesmo vale para as
fazendas Canabrava e Batinga.

Os totais são os da planilha, sem recalcular. Pode rodar de novo sem
duplicar. Para desfazer: delete from public.planejamento where importado is not null;

Precisa do openpyxl (pip install openpyxl).
"""
import sys
import unicodedata
import uuid
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

RAIZ = Path(__file__).resolve().parent.parent
PASTA = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
NS = uuid.UUID("5b3e8f0a-7c1d-4e2a-9f6b-0a1b2c3d4e5f")


def sid(grupo, n):
    return f"00000000-0000-4000-{grupo}-{n:012d}"


MILHO = sid("8000", 1)
# Bloco da planilha → (fazenda, talhão). Gameleira é área da Triunfo/Juerana.
FAZENDAS = {
    "TRIUNFO/JUERANA": ("Triunfo / Juerana", None),
    "GAMELEIRA": ("Triunfo / Juerana", sid("8200", 13)),
    "AGUAS CLARAS": ("Águas Claras", None),
    "SAO RAIMUNDO": ("São Raimundo", None),
    "CANABRAVA": ("Canabrava", None),
    "BATINGA": ("Batinga", None),
}
# As 8 colunas de fases da planilha, na ordem.
FASES = [
    "Pré-plantio / dessecação", "Plantio (adubação, TS, semente)", "Herbicida", "1ª cobertura",
    "1ª pulverização", "2ª cobertura", "2ª pulverização", "3ª pulverização",
]
ITENS = {"OPERACIONAL": ("Operacional", "Operacional"), "COLHEITA": ("Colheita", "Colheita"),
         "ARRENDAMENTO": ("Arrendamento", "Arrendamento"), "JUROS": ("Juros", "Juros / financiamento")}

# Nome na planilha → nome no cadastro (junta grafias diferentes do mesmo produto).
APELIDOS = {"AIRTUCK": "AIRTRUCK", "3707": "AS 3707", "BIOCROP10": "BIOCROP 10", "MAGNÉSIO": "MAGNESIO",
            "MANGANÊS": "MANGANES", "BORÔ LÍQUIDO": "BORO LIQUIDO", "BORÔ 13": "BORO 13"}
# Tipo e unidade do produto novo. Adubo em tonelada (a planilha de 2025 usa t/ha
# com preço por tonelada); o que não está aqui entra como "Outro", em litro.
TIPOS = {
    "adubo": (["KCL", "08-40-08", "SULFATO", "UREIA", "MAP"], "t"),
    "semente": (["AS 3707", "AS 9606", "AS 1868", "AS 1596"], "saco"),
    "herbicida": (["GLIFOSATO", "CRUCIAL", "2,4 D", "SONDA", "AURORA", "FLUMIZYN"], "L"),
    "inseticida": (["METOMIL", "CLORFENAPIR", "FIPRONIL", "VORAX"], "L"),
    "fungicida": (["CERCOBIN", "TIOFANATO", "PRIORI EXTRA", "TEBURAZ"], "L"),
    "foliar": (["MAGNESIO", "ZINCO", "MANGANES", "BORO LIQUIDO", "BORO 13", "COMONI", "POTASSIO",
                "MAP PURIFICADO", "NITROGENIO", "N 400", "MAGNUM", "ZINFLEX", "NORTHOFOS", "SPERTO", "KTOP"], "L"),
}
TIPO_DE = {p: (t, u) for t, (lista, u) in TIPOS.items() for p in lista}

avisos = []
sql = []
produtos = {}


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def nome_produto(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = " ".join(str(v).split()).upper()
    return APELIDOS.get(s, s)


def lit(v):
    if v is None:
        return "null"
    if isinstance(v, (int, float)):
        return repr(round(v, 6))
    return "'" + str(v).replace("'", "''") + "'"


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


# Compara nomes no banco sem acento e sem diferença de maiúscula.
CHAVE_SQL = "translate(upper(trim({c})), 'ÁÀÂÃÉÊÍÓÔÕÚÜÇ', 'AAAAEEIOOOUUC')"


def ref_por_nome(tabela, nome):
    return f"(select id from public.{tabela} where {CHAVE_SQL.format(c='nome')} = {lit(sem_acento(nome).upper())} order by criado_em limit 1)"


def item(origem, chave, reg):
    reg = {"id": str(uuid.uuid5(NS, f"PLANEJAMENTO/{origem}#{chave}")), **reg, "importado": f"PLANEJAMENTO/{origem}"}
    cols, vals = [], []
    for k, v in reg.items():
        cols.append(k)
        vals.append(v[1:] if isinstance(v, str) and v.startswith("=(") else lit(v))
    sql.append(f"insert into public.planejamento ({', '.join(cols)}) values ({', '.join(vals)}) on conflict (id) do nothing;")


def ler(arq, aba, safra, c0, c_res):
    achados = sorted(PASTA.glob(f"*{arq}*.xlsx"))
    if not achados:
        avisos.append(f"Planilha {arq} não está na pasta — ficou de fora")
        return []
    ws = openpyxl.load_workbook(achados[0], data_only=True)[aba]
    origem = f"{arq}/{aba.strip()}"
    resumo_blocos = []
    for cab in (r for r in range(1, ws.max_row + 1) if str(ws.cell(r, c0 + 1).value or "").strip() == "QTD/HA"):
        tit = next(r for r in range(cab, 0, -1) if str(ws.cell(r, c0).value or "").upper().startswith("FAZENDA"))
        bloco = str(ws.cell(tit, c0).value).strip()[8:].strip().upper()
        fazenda, talhao = FAZENDAS[bloco]
        ha = 0.0
        for r in range(tit, tit + 3):
            for c in range(c0 + 30, c0 + 45):
                if str(ws.cell(r, c).value or "").strip() == "HECTARES":
                    ha = num(ws.cell(r, c + 1).value) or num(ws.cell(r + 1, c + 1).value)
        fim = next(r for r in range(cab + 1, ws.max_row + 1) if str(ws.cell(r, c0).value or "").strip() == "TOTAIS")
        base = {"safra": safra, "cultura_id": MILHO, "fazenda_id": "=" + ref_por_nome("fazendas", fazenda), "talhao_id": talhao, "hectares": ha}
        soma = 0.0
        for g, fase in enumerate(FASES):
            cg = c0 + 5 * g
            for r in range(cab + 1, fim):
                p, dose, qtd, preco, total = (ws.cell(r, cg + i).value for i in range(5))
                if p is None or str(p).strip() in ("", "TOTAIS") or not (num(dose) or num(total)):
                    continue
                prod = nome_produto(p)
                tipo, unidade = TIPO_DE.get(prod, ("outro", "L"))
                produtos.setdefault(prod, (tipo, unidade))
                dose, qtd = num(dose), num(qtd)
                if tipo == "adubo" and dose >= 5:  # 2026 em kg/ha: o cadastro do adubo é em t
                    dose, qtd = dose / 1000, qtd / 1000
                if qtd and ha and abs(dose * ha - qtd) > 0.01:
                    avisos.append(f"{origem} › {bloco} › {prod} ({fase}): quantidade total {qtd:g} não é dose {dose:g} × {ha:g} ha = {dose * ha:g} (célula {openpyxl.utils.get_column_letter(cg + 2)}{r})")
                item(origem, f"{bloco}/{openpyxl.utils.get_column_letter(cg)}{r}", {
                    **base, "fase": fase, "insumo_id": "=" + ref_por_nome("insumos", prod), "descricao": None,
                    "dose_ha": dose, "quantidade_total": qtd or dose * ha, "preco_unitario": num(preco) or None, "total": num(total),
                })
                soma += num(total)
        resumo = {}
        for r in range(tit, fim + 1):
            k = str(ws.cell(r, c_res).value or "").strip().upper()
            if k in ITENS or k in ("TOTAL", "ADUBACAO", "SEMENTES", "QUIMICOS"):
                resumo.setdefault(k, num(ws.cell(r, c_res + 1).value))
        for k, (descricao, fase) in ITENS.items():
            v = resumo.get(k, 0)
            if not v:
                continue
            item(origem, f"{bloco}/{k}", {
                **base, "fase": fase, "insumo_id": None, "descricao": descricao,
                "dose_ha": 1, "quantidade_total": ha, "preco_unitario": round(v / ha, 4) if ha else v, "total": v,
            })
            soma += v
        dif = soma - resumo.get("TOTAL", 0)
        if abs(dif) >= 0.01:
            avisos.append(f"{origem} › {bloco}: soma dos itens R$ {soma:,.2f} × total da planilha R$ {resumo.get('TOTAL', 0):,.2f} (diferença R$ {dif:,.2f})")
        resumo_blocos.append((safra, bloco, ha, soma))
    return resumo_blocos


blocos = []
blocos += ler("MILHO_2025", "PLANEJAMENTO", "Milho 2025", ci("C"), ci("AR"))
blocos += ler("MILHO_2026", "PLANEJAMENTO", "Milho 2026", ci("C"), ci("AR"))
blocos += ler("MILHO_2025", "SOCIEDADE GILBERTO ", "Milho 2025 (sociedade Gilberto)", ci("B"), ci("AQ"))

cadastros = []
for nome in ("Canabrava", "Batinga"):
    cadastros.append(
        f"insert into public.fazendas (id, nome, posse, ativo, observacao, importado) select {lit(str(uuid.uuid5(NS, 'fazenda/' + nome)))}, {lit(nome)}, 'propria', true, "
        f"'Criada pela importação do planejamento do milho — confira posse e área', 'PLANEJAMENTO' "
        f"where not exists (select 1 from public.fazendas where {CHAVE_SQL.format(c='nome')} = {lit(nome.upper())});"
    )
for prod, (tipo, unidade) in sorted(produtos.items()):
    cadastros.append(
        f"insert into public.insumos (id, nome, tipo, unidade, ativo, observacao, importado) select {lit(str(uuid.uuid5(NS, 'insumo/' + prod)))}, {lit(prod)}, {lit(tipo)}, {lit(unidade)}, true, "
        f"'Criado pela importação do planejamento do milho — confira tipo e unidade', 'PLANEJAMENTO' "
        f"where not exists (select 1 from public.insumos where {CHAVE_SQL.format(c='nome')} = {lit(sem_acento(prod).upper())});"
    )

schema = (RAIZ / "supabase/schema.sql").read_text()
total = sum(b[3] for b in blocos)
saida = RAIZ / "supabase/importacao-planejamento.sql"
saida.write_text(f"""-- ════════════════════════════════════════════════════════════════════════
-- Planejamento do milho — GERADO por scripts/importar-planejamento.py
-- {len(sql)} itens, {len(blocos)} blocos de fazenda, R$ {total:,.2f} planejados; {len(produtos)} produtos
-- (os que não existirem no cadastro são criados).
--
-- Cole tudo no SQL Editor do Supabase e clique em Run. Pode rodar de novo:
-- não duplica. Para desfazer:
--   delete from public.planejamento where importado is not null;
-- ════════════════════════════════════════════════════════════════════════

begin;

{schema}
-- Marca do que veio das planilhas.
alter table public.planejamento add column if not exists importado text;
alter table public.insumos add column if not exists importado text;
alter table public.fazendas add column if not exists importado text;

-- Fazendas e produtos que faltarem no cadastro.
{chr(10).join(cadastros)}

-- Itens do planejamento.
{chr(10).join(sql)}

commit;
""")
print(f"{saida.relative_to(RAIZ)}: {len(sql)} itens, R$ {total:,.2f}")
for safra, bloco, ha, soma in blocos:
    print(f"  {safra:34} {bloco:16} {ha:5g} ha  R$ {soma:12,.2f}")
print(f"\n{len(avisos)} avisos:")
for a in avisos:
    print(" -", a)
